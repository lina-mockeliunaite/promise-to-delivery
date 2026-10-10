"""Document adapters: one function per format, each turning an uploaded file into canonical text plus a list of problems.

Adapters only produce text. Extraction receives that text through the existing path (extraction_cache.get_or_extract ->
extract.attempt_extraction(..., text=...)); nothing in the evaluated pipeline changes. Markdown and text delegate to
ledger_import.markdown_adapter, so their canonical text, hashes and adapter name/version are exactly what they were.

PDF: text layer only, with a "[Page N]" marker line before each page's text. Excel and PowerPoint (9 Oct):
docs/SPEC_xlsx_pptx_canonical_text.md. DOCX: body paragraphs and tables in order
(headings are their text; table rows are flattened with " | " inside the canonical text only; headers and footers are
not read). Neither format is OCR'd. Nothing here reads a file path or touches the database except the notes helpers.
"""

import io
import json
import re
import zipfile
from pathlib import Path

import check_quotes
import ledger_import

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_PDF_PAGES = 200
MAX_CANONICAL_CHARS = 1_000_000
MAX_DOCX_UNPACKED = 50 * 1024 * 1024
NOTES_SCHEMA_PATH = Path(__file__).resolve().parent / "adapter_notes_schema.sql"

PDF_ADAPTER_NAME, DOCX_ADAPTER_NAME = "pdf-text", "docx"
XLSX_ADAPTER_NAME, PPTX_ADAPTER_NAME = "xlsx", "pptx"
FORMAT_LABELS = {".md": "Markdown", ".txt": "Text", ".pdf": "PDF", ".docx": "Word", ".xlsx": "Excel", ".pptx": "PowerPoint"}
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
PAGE_MARKER = re.compile(r"^\[Page (\d+)\]$", re.M)
NO_TEXT_PAGE = "Page {n} has no readable text (scanned or image-only; OCR is not supported in this build)."
NO_TEXT_PDF = "No page in this PDF has readable text (scanned or image-only; OCR is not supported in this build)."
SUPPORTED = "This build reads PDF, Word (.docx), Excel (.xlsx), PowerPoint (.pptx), Markdown and text files."

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W, "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "v": "urn:schemas-microsoft-com:vml", "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006"}


class AdapterError(Exception):
    """A document the build cannot read. The message is plain language for the person uploading it."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.status = status


def extension(filename: str) -> str:
    name = (filename or "").replace("\\", "/").rsplit("/", 1)[-1].lower()
    return "." + name.rsplit(".", 1)[-1] if "." in name else ""


def format_label(adapter_name: str, filename: str = "") -> str:
    """The format a person recognises, from what the ledger recorded."""
    if adapter_name == PDF_ADAPTER_NAME:
        return "PDF"
    if adapter_name == DOCX_ADAPTER_NAME:
        return "Word"
    if adapter_name == XLSX_ADAPTER_NAME:
        return "Excel"
    if adapter_name == PPTX_ADAPTER_NAME:
        return "PowerPoint"
    return "Text" if extension(filename) == ".txt" else "Markdown"


def adapt(filename: str, raw: bytes) -> dict:
    """{canonical_text, location_map (JSON text), adapter_name, adapter_version, format, problems} or AdapterError."""
    ext = extension(filename)
    if ext == ".doc":
        raise AdapterError("Older Word files (.doc) are not supported. Save the document as .docx and upload that.")
    if ext in (".xls", ".xlsm", ".xlsb"):
        raise AdapterError("This build reads Excel files saved as .xlsx (no macros). Save the workbook as .xlsx and upload that.")
    if ext in (".ppt", ".pptm", ".ppsx", ".pps"):
        raise AdapterError("This build reads PowerPoint files saved as .pptx (no macros). Save the deck as .pptx and upload that.")
    if ext not in FORMAT_LABELS:
        raise AdapterError(SUPPORTED)
    if not raw:
        raise AdapterError("The file is empty.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise AdapterError("The file is larger than 10 MB.", 413)
    if ext in (".md", ".txt"):
        try:
            out = ledger_import.markdown_adapter(raw)
        except UnicodeDecodeError:
            raise AdapterError("This text file is not UTF-8 text. Save it as UTF-8 and upload it again.") from None
        if not out["canonical_text"].strip():
            raise AdapterError("The document is empty.")
        return {**out, "format": FORMAT_LABELS[ext], "problems": []}
    out = {".pdf": _adapt_pdf, ".docx": _adapt_docx, ".xlsx": _adapt_xlsx, ".pptx": _adapt_pptx}[ext](raw)
    if len(out["canonical_text"]) > MAX_CANONICAL_CHARS:
        raise AdapterError("This document has too much text to read in this build (over 1,000,000 characters).")
    return out


# --- PDF ------------------------------------------------------------------------------------------------------------

def _clean_page(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    lines = [line.rstrip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _adapt_pdf(raw: bytes) -> dict:
    import pypdf
    if b"%PDF-" not in raw[:1024]:
        raise AdapterError("This file is named .pdf but is not a PDF.")
    try:
        reader = pypdf.PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise AdapterError("This PDF is password-protected. Remove the password and upload it again.")
        pages = reader.pages
        count = len(pages)
        if count == 0:
            raise AdapterError("This PDF has no pages.")
        if count > MAX_PDF_PAGES:
            raise AdapterError(f"This PDF has {count} pages; this build reads up to {MAX_PDF_PAGES}.")
        texts = [_clean_page(pages[i].extract_text() or "") for i in range(count)]
    except AdapterError:
        raise
    except Exception:  # pypdf raises many kinds for damaged files; none of them is worth showing
        raise AdapterError("This PDF could not be read. It may be damaged; export it again and upload it.") from None
    problems, parts, location, offset = [], [], [], 0
    for number, text in enumerate(texts, start=1):
        if not text:
            problems.append(NO_TEXT_PAGE.format(n=number))
            continue
        block = f"[Page {number}]\n{text}"
        if parts:
            offset += 2  # the blank line between pages
        location.append({"page": number, "start": offset, "end": offset + len(block)})
        parts.append(block)
        offset += len(block)
    if not parts:
        raise AdapterError(NO_TEXT_PDF)
    return {"canonical_text": "\n\n".join(parts), "location_map": json.dumps(location),
            "adapter_name": PDF_ADAPTER_NAME, "adapter_version": f"1 (pypdf {pypdf.__version__})",
            "format": "PDF", "problems": problems}


# --- DOCX -----------------------------------------------------------------------------------------------------------

def _q(tag: str) -> str:
    prefix, name = tag.split(":")
    return f"{{{NS[prefix]}}}{name}"


def _in_fallback(el) -> bool:
    return any(a.tag == _q("mc:Fallback") for a in el.iterancestors())


def _paragraph_text(p) -> str:
    out = []
    for el in p.iter():
        if el.tag == _q("w:t"):
            if not any(a.tag == _q("w:txbxContent") for a in el.iterancestors()):
                out.append(el.text or "")
        elif el.tag == _q("w:tab"):
            out.append("\t")
        elif el.tag in (_q("w:br"), _q("w:cr")) and not any(a.tag == _q("w:txbxContent") for a in el.iterancestors()):
            out.append("\n")
    return "".join(out).strip()


def _blocks(container):
    """Paragraph strings and table-row strings of a body or cell, in document order."""
    for child in container.iterchildren():
        if child.tag == _q("w:p"):
            text = _paragraph_text(child)
            if text:
                yield text
        elif child.tag == _q("w:tbl"):
            rows = []
            for tr in child.iterchildren(_q("w:tr")):
                cells = []
                for tc in tr.iterchildren(_q("w:tc")):
                    cells.append(" ".join(" ".join(b.split()) for b in _blocks(tc)))
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                yield "\n".join(rows)
        elif child.tag == _q("w:sdt"):
            for content in child.iterchildren(_q("w:sdtContent")):
                yield from _blocks(content)


def _plural(n: int, one: str, many: str) -> str:
    return f"{n} {one if n == 1 else many}"


def _adapt_docx(raw: bytes) -> dict:
    import docx
    from lxml import etree
    if raw.startswith(OLE_MAGIC):
        raise AdapterError("That is an older or password-protected Word file. Save it as .docx without a password and upload that.")
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise AdapterError("This file is named .docx but is not a Word (.docx) document.")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            if sum(i.file_size for i in infos) > MAX_DOCX_UNPACKED:
                raise AdapterError("This Word document is too large once unpacked to read in this build.")
            names = {i.filename for i in infos}
            if "word/document.xml" not in names:
                raise AdapterError("This file is named .docx but is not a Word (.docx) document.")
            footnotes = comments = 0
            if "word/footnotes.xml" in names:
                root = etree.fromstring(zf.read("word/footnotes.xml"), etree.XMLParser(resolve_entities=False, no_network=True))
                footnotes = sum(1 for f in root.iter(_q("w:footnote"))
                                if f.get(_q("w:type")) not in ("separator", "continuationSeparator", "continuationNotice")
                                and "".join(f.itertext()).strip())
            if "word/comments.xml" in names:
                root = etree.fromstring(zf.read("word/comments.xml"), etree.XMLParser(resolve_entities=False, no_network=True))
                comments = sum(1 for _ in root.iter(_q("w:comment")))
        document = docx.Document(io.BytesIO(raw))
    except AdapterError:
        raise
    except Exception:
        raise AdapterError("This Word document could not be read. It may be damaged; save it again and upload it.") from None
    body = document.element.body
    canonical = "\n\n".join(_blocks(body)).replace("\r\n", "\n").replace("\r", "\n")
    if not canonical.strip():
        raise AdapterError("This Word document has no text to read.")
    images = sum(1 for b in body.iter(_q("a:blip")) if not _in_fallback(b)) + sum(
        1 for i in body.iter(_q("v:imagedata")) if not _in_fallback(i))
    objects = sum(1 for o in body.iter(_q("w:object")) if not _in_fallback(o))
    boxes = sum(1 for t in body.iter(_q("w:txbxContent")) if not _in_fallback(t))
    changes = sum(1 for el in body.iter(_q("w:ins"), _q("w:del")) if not _in_fallback(el))
    problems = []
    if images:
        problems.append(f"{_plural(images, 'image was', 'images were')} not read (pictures are not supported in this build).")
    if objects:
        problems.append(f"{_plural(objects, 'embedded object was', 'embedded objects were')} not read.")
    if boxes:
        problems.append(f"{_plural(boxes, 'text box was', 'text boxes were')} not read.")
    if footnotes:
        problems.append("Footnotes are not read.")
    if comments:
        problems.append("Comments are not read.")
    if changes:
        problems.append(f"The document has {_plural(changes, 'tracked change', 'tracked changes')}; "
                        "the text read is the version with every change accepted.")
    return {"canonical_text": canonical, "location_map": ledger_import.markdown_adapter(canonical.encode("utf-8"))["location_map"],
            "adapter_name": DOCX_ADAPTER_NAME, "adapter_version": f"1 (python-docx {_docx_version()})",
            "format": "Word", "problems": problems}


def _docx_version() -> str:
    import importlib.metadata
    return importlib.metadata.version("python-docx")


# --- Quotes trace back ----------------------------------------------------------------------------------------------

def _flatten(canonical_text: str, pages: list):
    """(text, offsets): the canonical text with page markers removed and whitespace collapsed, and for each character
    of that text its offset in the canonical text."""
    skip = set()
    for entry in pages:
        marker = f"[Page {entry['page']}]"
        if canonical_text.startswith(marker, entry["start"]):
            skip.update(range(entry["start"], entry["start"] + len(marker)))
    out, offsets, pending_space = [], [], False
    for i, ch in enumerate(canonical_text):
        if i in skip:
            pending_space = True
        elif ch.isspace():
            pending_space = True
        else:
            if pending_space and out:
                out.append(" ")
                offsets.append(i)
            pending_space = False
            out.append(ch)
            offsets.append(i)
    return "".join(out), offsets


def _pages(location_map) -> list:
    try:
        data = json.loads(location_map) if isinstance(location_map, str) else location_map
    except ValueError:
        return []
    return [e for e in data or [] if isinstance(e, dict) and "page" in e]


def quote_in_text(canonical_text: str, quote: str) -> bool:
    """True if the quote appears word for word in the canonical text (whitespace and page markers ignored)."""
    flat, _ = _flatten(canonical_text, _pages(json.dumps(_marker_pages(canonical_text))))
    return check_quotes.normalise(quote) in flat


def _marker_pages(canonical_text: str) -> list:
    """Page entries recovered from the markers themselves, for text whose location map is not at hand."""
    found = [(int(m.group(1)), m.start()) for m in PAGE_MARKER.finditer(canonical_text)]
    return [{"page": n, "start": s, "end": (found[k + 1][1] if k + 1 < len(found) else len(canonical_text))}
            for k, (n, s) in enumerate(found)]


def locate_page(canonical_text: str, location_map, quote: str):
    """The page a quote starts on, or None if it is not found or appears on more than one page."""
    pages = _pages(location_map)
    if not pages:
        return None
    flat, offsets = _flatten(canonical_text, pages)
    needle = check_quotes.normalise(quote)
    if not needle:
        return None
    found, at = set(), flat.find(needle)
    while at != -1:
        offset = offsets[at]
        page = next((e["page"] for e in pages if e["start"] <= offset < e["end"] + 2), None)
        found.add(page)
        at = flat.find(needle, at + 1)
    return found.pop() if len(found) == 1 else None


# --- Excel (.xlsx) ----------------------------------------------------------------------------------------------------
# Spec: docs/SPEC_xlsx_pptx_canonical_text.md. Values only (a formula counts by its stored value), hidden content read
# and flagged, one row per line with " | " between cells, a "[Sheet: name]" marker before each sheet.

NO_VALUE = "[no stored value]"
MAX_XLSX_CELLS = 200_000
MAX_PPTX_SLIDES = 300
MAX_OOXML_UNPACKED = 50 * 1024 * 1024
SHEET_MARKER = re.compile(r"^\[Sheet: .*\]$", re.M)
SLIDE_MARKER = re.compile(r"^\[Slide (\d+)(?: \(hidden\))?\]$", re.M)


def _ooxml_zip(raw: bytes, kind: str, required: str) -> None:
    """The checks every Office Open XML upload gets before a library opens it."""
    article = "an" if kind[0] in "AEIOU" else "a"
    if raw.startswith(OLE_MAGIC):
        raise AdapterError(f"That is an older or password-protected {kind} file. Save it as "
                           f"{'.xlsx' if kind == 'Excel' else '.pptx'} without a password and upload that.")
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise AdapterError(f"This file is not {article} {kind} document.")
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        infos = zf.infolist()
        if sum(i.file_size for i in infos) > MAX_OOXML_UNPACKED:
            raise AdapterError(f"This {kind} file is too large once unpacked to read in this build.")
        names = {i.filename for i in infos}
        if required not in names:
            raise AdapterError(f"This file is not {article} {kind} document.")
        if any(n.endswith("vbaProject.bin") for n in names):
            raise AdapterError(f"This {kind} file contains macros. Save it without macros and upload that.")


def _cell_text(value, data_type) -> str:
    import datetime
    if value is None:
        return ""
    if data_type == "e":
        return NO_VALUE
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime.datetime):
        return value.date().isoformat() if value.time() == datetime.time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return " ".join(str(value).split())


def _adapt_xlsx(raw: bytes) -> dict:
    import openpyxl
    from openpyxl.utils import get_column_letter
    _ooxml_zip(raw, "Excel", "xl/workbook.xml")
    try:
        formulas = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
        values = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    except AdapterError:
        raise
    except Exception:
        raise AdapterError("This Excel file could not be read. It may be damaged; save it again and upload it.") from None
    problems, parts, location, offset, cells_seen = [], [], [], 0, 0
    missing, hidden_rows, hidden_cols, comments = 0, 0, 0, 0
    charts = images = 0
    for sheet in formulas.worksheets:
        vs = values[sheet.title]
        hidden_sheet = sheet.sheet_state != "visible"
        if hidden_sheet:
            problems.append(f"Sheet “{sheet.title}” is hidden; its cells are read and marked hidden.")
        charts += len(getattr(sheet, "_charts", []))
        images += len(getattr(sheet, "_images", []))
        col_hidden = {k for k, d in sheet.column_dimensions.items() if d.hidden}
        hidden_cols += len(col_hidden)
        lines, entries = [], []
        for row in sheet.iter_rows():
            texts, first, last, any_hidden_col = [], None, None, False
            for cell in row:
                if cell.value is None:
                    texts.append("")
                    continue
                cells_seen += 1
                if cells_seen > MAX_XLSX_CELLS:
                    raise AdapterError(f"This workbook has more than {MAX_XLSX_CELLS:,} filled cells; this build reads up to that.")
                if cell.comment is not None:
                    comments += 1
                v = vs[cell.coordinate]
                if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    text = NO_VALUE if v.value is None else _cell_text(v.value, v.data_type)
                else:
                    text = _cell_text(v.value, v.data_type)
                if text == NO_VALUE:
                    missing += 1
                if text:
                    first = first or cell.coordinate
                    last = cell.coordinate
                    if get_column_letter(cell.column) in col_hidden:
                        any_hidden_col = True
                texts.append(text)
            while texts and not texts[-1]:
                texts.pop()
            if not any(texts):
                continue
            r = row[0].row
            row_hidden = bool(sheet.row_dimensions[r].hidden) if r in sheet.row_dimensions else False
            hidden_rows += row_hidden
            lines.append(" | ".join(texts))
            entries.append({"sheet": sheet.title, "row": r, "cells": f"{first}:{last}",
                            "hidden": "sheet" if hidden_sheet else "row" if row_hidden else "column" if any_hidden_col else None})
        if not lines:
            continue
        marker = f"[Sheet: {sheet.title}{' (hidden)' if hidden_sheet else ''}]"
        if parts:
            offset += 2
        offset += len(marker) + 1
        for line, entry in zip(lines, entries):
            location.append({**entry, "start": offset, "end": offset + len(line)})
            offset += len(line) + 1
        offset -= 1
        parts.append(marker + "\n" + "\n".join(lines))
    if formulas.chartsheets:
        problems.append(f"{_plural(len(formulas.chartsheets), 'chart sheet was', 'chart sheets were')} not read.")
    if not parts:
        raise AdapterError("This Excel file has no cell text to read.")
    if hidden_rows:
        problems.append(f"{_plural(hidden_rows, 'hidden row is', 'hidden rows are')} read and marked hidden.")
    if hidden_cols:
        problems.append(f"{_plural(hidden_cols, 'hidden column is', 'hidden columns are')} read and marked hidden.")
    if missing:
        problems.append(f"{_plural(missing, 'formula cell has', 'formula cells have')} no stored value or an error; "
                        f"shown as {NO_VALUE}, never recalculated. Save the file in Excel and upload it again to store values.")
    if comments:
        problems.append("Cell comments are not read.")
    if charts or images:
        problems.append("Charts and pictures are not read.")
    return {"canonical_text": "\n\n".join(parts), "location_map": json.dumps(location),
            "adapter_name": XLSX_ADAPTER_NAME, "adapter_version": f"1 (openpyxl {openpyxl.__version__})",
            "format": "Excel", "problems": problems}


# --- PowerPoint (.pptx) -----------------------------------------------------------------------------------------------

def _shape_lines(shape):
    """Text lines of one shape, groups opened, tables one row per line."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for inner in shape.shapes:
            yield from _shape_lines(inner)
        return
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                yield " | ".join(cells)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            text = " ".join("".join(r.text for r in p.runs).split()) or " ".join((p.text or "").split())
            if text:
                yield text


def _unread(shape) -> int:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return sum(_unread(s) for s in shape.shapes)
    if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.MEDIA,
                            MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
        return 1
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return 1
    if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and shape.__class__.__name__ == "PlaceholderPicture":
        return 1
    return 1 if shape.__class__.__name__ == "GraphicFrame" and not getattr(shape, "has_table", False) else 0


def _adapt_pptx(raw: bytes) -> dict:
    import pptx
    _ooxml_zip(raw, "PowerPoint", "ppt/presentation.xml")
    try:
        deck = pptx.Presentation(io.BytesIO(raw))
        slides = list(deck.slides)
    except Exception:
        raise AdapterError("This PowerPoint file could not be read. It may be damaged; save it again and upload it.") from None
    if not slides:
        raise AdapterError("This PowerPoint file has no slides.")
    if len(slides) > MAX_PPTX_SLIDES:
        raise AdapterError(f"This deck has {len(slides)} slides; this build reads up to {MAX_PPTX_SLIDES}.")
    problems, parts, location, offset = [], [], [], 0
    unread = notes = 0
    hidden_list, empty_list = [], []
    for number, slide in enumerate(slides, start=1):
        hidden = slide._element.get("show") == "0"
        shapes = list(slide.shapes)
        title = slide.shapes.title
        ordered = ([title] if title is not None else []) + [s for s in shapes if title is None or s.shape_id != title.shape_id]
        lines = [line for s in ordered for line in _shape_lines(s)]
        unread += sum(_unread(s) for s in shapes)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None and slide.notes_slide.notes_text_frame.text.strip():
            notes += 1
        if hidden:
            hidden_list.append(number)
        if not lines:
            empty_list.append(number)
            continue
        block = f"[Slide {number}{' (hidden)' if hidden else ''}]\n" + "\n".join(lines)
        if parts:
            offset += 2
        location.append({"slide": number, "hidden": hidden, "start": offset, "end": offset + len(block)})
        parts.append(block)
        offset += len(block)
    if not parts:
        raise AdapterError("This deck has no slide text to read (pictures and charts are not read in this build).")
    if hidden_list:
        problems.append(f"Hidden {'slide' if len(hidden_list) == 1 else 'slides'} {', '.join(map(str, hidden_list))}: read and marked hidden.")
    if empty_list:
        problems.append(f"No readable text on {'slide' if len(empty_list) == 1 else 'slides'} {', '.join(map(str, empty_list))}.")
    if unread:
        problems.append(f"{_plural(unread, 'picture, chart or object was', 'pictures, charts or objects were')} not read.")
    if notes:
        problems.append("Speaker notes are not read.")
    return {"canonical_text": "\n\n".join(parts), "location_map": json.dumps(location),
            "adapter_name": PPTX_ADAPTER_NAME, "adapter_version": f"1 (python-pptx {pptx.__version__})",
            "format": "PowerPoint", "problems": problems}


# --- Where a quote sits, for any format ------------------------------------------------------------------------------

def locate(canonical_text: str, location_map, quote: str, adapter_name: str):
    """A plain location for a quote: 'page 3', 'sheet Approvals, row 5', 'slide 4'; None when unknown or ambiguous."""
    if adapter_name == PDF_ADAPTER_NAME:
        page = locate_page(canonical_text, location_map, quote)
        return f"page {page}" if page else None
    if adapter_name not in (XLSX_ADAPTER_NAME, PPTX_ADAPTER_NAME):
        return None
    try:
        entries = json.loads(location_map) if isinstance(location_map, str) else (location_map or [])
    except ValueError:
        return None
    needle = check_quotes.normalise(quote)
    if not needle:
        return None
    found = set()
    for e in entries:
        if not isinstance(e, dict):
            continue
        piece = check_quotes.normalise(canonical_text[e.get("start", 0):e.get("end", 0)])
        if needle in piece:
            if "slide" in e:
                found.add(f"slide {e['slide']}" + (" (hidden)" if e.get("hidden") else ""))
            elif "sheet" in e:
                found.add(f"sheet {e['sheet']}, row {e['row']}" + (f" (hidden {e['hidden']})" if e.get("hidden") else ""))
    return found.pop() if len(found) == 1 else None


# --- Notes (problems) stored with a source version ------------------------------------------------------------------

def ensure_notes_schema(conn) -> None:
    have = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE name LIKE 'source_version_notes%'")}
    if {"source_version_notes", "source_version_notes_immutable_u", "source_version_notes_immutable_d"} <= have:
        return
    conn.executescript(NOTES_SCHEMA_PATH.read_text(encoding="utf-8"))
    conn.commit()


def save_notes(conn, source_version_id: int, problems: list) -> None:
    """Store a version's problems. The caller owns the transaction; the table must exist (ensure_notes_schema)."""
    for n, message in enumerate(problems, start=1):
        conn.execute("INSERT INTO source_version_notes (source_version_id, ordinal, message) VALUES (?, ?, ?)",
                     (source_version_id, n, message))


def notes_for(conn, version_ids: list) -> dict:
    """{source_version_id: [message, ...]}. Empty if no document with problems was ever uploaded to this database."""
    if not version_ids or not conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'source_version_notes'").fetchone():
        return {}
    marks = ",".join("?" * len(version_ids))
    out = {}
    for vid, message in conn.execute(
            f"SELECT source_version_id, message FROM source_version_notes WHERE source_version_id IN ({marks})"
            " ORDER BY source_version_id, ordinal", list(version_ids)):
        out.setdefault(vid, []).append(message)
    return out
