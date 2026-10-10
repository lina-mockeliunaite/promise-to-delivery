"""Excel and PowerPoint adapters (built 9 Oct; spec docs/SPEC_xlsx_pptx_canonical_text.md).

Fixtures are generated in memory from development text only (the Harbour Bank scenario pricing note); nothing under
data/ is written and nothing under data/coral_pay/ is read. No model, no network.
"""

import io
import json
import os
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
import pptx
from fastapi.testclient import TestClient
from pptx.util import Inches

import adapters
import api
import check_quotes
import config

H = {"X-Requested-With": "deal-workspace"}


def note_rows():
    """The approval table of the HB-05 v2 scenario note, as rows of cells."""
    text = config.scenario_path("HB-05_v2_named_exception.md").read_text(encoding="utf-8")
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in text.splitlines()
            if line.startswith("| CAP-") or line.startswith("| Capability")]
    return rows


def xlsx_bytes(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def pricing_workbook(wb):
    ws = wb.active
    ws.title = "Approvals"
    ws.append(["INTERNAL — Pricing and services note: Harbour Bank"])
    ws.append(["Date", "15 October 2026"])
    ws.append([])
    for row in note_rows():
        ws.append(row)


def with_cached_value(raw: bytes, cell: str, value: str) -> bytes:
    """Give a formula cell a stored value, as Excel would on save (openpyxl writes formulas without one)."""
    src = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                xml = data.decode("utf-8")
                xml = re.sub(rf'(<c r="{cell}"[^>]*>)(<f>[^<]*</f>)(<v\s*/>|<v></v>)?', rf"\1\2<v>{value}</v>", xml)
                data = xml.encode("utf-8")
            dst.writestr(info, data)
    return out.getvalue()


class TestExcel(unittest.TestCase):
    def adapt(self, raw, name="note.xlsx"):
        return adapters.adapt(name, raw)

    def test_one_row_per_line_with_the_approval_in_the_last_cell(self):
        out = self.adapt(xlsx_bytes(pricing_workbook))
        self.assertEqual(out["format"], "Excel")
        self.assertEqual(out["adapter_name"], "xlsx")
        lines = out["canonical_text"].splitlines()
        self.assertEqual(lines[0], "[Sheet: Approvals]")
        polygon = next(l for l in lines if "Polygon, real-time" in l)
        self.assertTrue(polygon.endswith("subject to the beta terms in Annex B"))
        self.assertEqual(polygon.count(" | "), 3)
        self.assertEqual(out["problems"], [])

    def test_the_location_map_points_at_each_row(self):
        out = self.adapt(xlsx_bytes(pricing_workbook))
        text, entries = out["canonical_text"], json.loads(out["location_map"])
        for e in entries:
            self.assertNotIn("\n", text[e["start"]:e["end"]])
            self.assertTrue(text[e["start"]:e["end"]].strip())
        self.assertEqual(adapters.locate(text, out["location_map"], "Named exception approved by Daniel Koh", "xlsx"),
                         "sheet Approvals, row 6")
        self.assertTrue(adapters.quote_in_text(text, "Named exception approved by Daniel Koh (Head of Product)"))

    def test_same_file_same_text(self):
        raw = xlsx_bytes(pricing_workbook)
        self.assertEqual(self.adapt(raw)["canonical_text"], self.adapt(raw)["canonical_text"])

    def test_hidden_sheets_rows_and_columns_are_read_and_flagged(self):
        def build(wb):
            pricing_workbook(wb)
            wb.active.row_dimensions[6].hidden = True
            secret = wb.create_sheet("Deal desk")
            secret.append(["CAP-024 VASP counterparty data exchange", "Named exception approved by Mei Tan (CRO)"])
            secret.sheet_state = "hidden"
            third = wb.create_sheet("Margins")
            third.append(["a", "b", "c"])
            third.column_dimensions["B"].hidden = True
        out = self.adapt(xlsx_bytes(build))
        self.assertIn("[Sheet: Deal desk (hidden)]", out["canonical_text"])
        self.assertIn("Named exception approved by Mei Tan", out["canonical_text"])  # read, not dropped
        hidden = {e["hidden"] for e in json.loads(out["location_map"])}
        self.assertEqual(hidden, {None, "row", "sheet", "column"})
        joined = " ".join(out["problems"])
        self.assertIn("Sheet “Deal desk” is hidden", joined)
        self.assertIn("1 hidden row is read", joined)
        self.assertIn("1 hidden column is read", joined)
        self.assertEqual(adapters.locate(out["canonical_text"], out["location_map"], "Named exception approved by Daniel Koh", "xlsx"),
                         "sheet Approvals, row 6 (hidden row)")

    def test_a_formula_counts_only_by_its_stored_value(self):
        def build(wb):
            ws = wb.active
            ws.title = "Limits"
            ws.append(["CAP-023 Payout ledger connector", 25000, "=B1*2"])
            ws.append(["Total", "=SUM(B1:B1)"])
        raw = xlsx_bytes(build)
        out = self.adapt(raw)
        self.assertIn(f"CAP-023 Payout ledger connector | 25000 | {adapters.NO_VALUE}", out["canonical_text"])
        self.assertNotIn("50000", out["canonical_text"])  # never recalculated
        self.assertNotIn("=B1", out["canonical_text"])     # never the formula
        self.assertIn("2 formula cells have no stored value", " ".join(out["problems"]))
        stored = self.adapt(with_cached_value(raw, "C1", "50000"))
        self.assertIn("CAP-023 Payout ledger connector | 25000 | 50000", stored["canonical_text"])
        self.assertIn("1 formula cell has no stored value", " ".join(stored["problems"]))

    def test_an_error_value_is_not_a_value(self):
        def build(wb):
            wb.active.append(["CAP-021", "=1/0"])
        # Excel stores an error result with t="e"; write the sheet that way
        raw = xlsx_bytes(build)
        src = zipfile.ZipFile(io.BytesIO(raw))
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename == "xl/worksheets/sheet1.xml":
                    data, n = re.subn(rb'<c r="B1"([^>]*)><f>([^<]*)</f>(<v\s*/>|<v></v>)?</c>', rb'<c r="B1"\1 t="e"><f>\2</f><v>#DIV/0!</v></c>', data)
                    self.assertEqual(n, 1)  # the error value really is in the file
                dst.writestr(info, data)
        out = self.adapt(buf.getvalue())
        self.assertIn(f"CAP-021 | {adapters.NO_VALUE}", out["canonical_text"])
        self.assertNotIn("#DIV/0!", out["canonical_text"])

    def test_numbers_dates_and_booleans_read_plainly(self):
        import datetime
        def build(wb):
            wb.active.append(["Limit", 40000.0, datetime.datetime(2026, 10, 14), True, 2.5])
        self.assertIn("Limit | 40000 | 2026-10-14 | TRUE | 2.5", self.adapt(xlsx_bytes(build))["canonical_text"])

    def test_refusals(self):
        raw = xlsx_bytes(pricing_workbook)
        for name, message in (("old.xls", ".xlsx"), ("macro.xlsm", "no macros")):
            with self.assertRaises(adapters.AdapterError) as ctx:
                adapters.adapt(name, raw)
            self.assertIn(message, str(ctx.exception))
        with self.assertRaises(adapters.AdapterError):
            adapters.adapt("fake.xlsx", b"not a zip at all")
        with self.assertRaises(adapters.AdapterError) as ctx:
            adapters.adapt("old.xlsx", adapters.OLE_MAGIC + b"\0" * 100)
        self.assertIn("password-protected", str(ctx.exception))
        empty = xlsx_bytes(lambda wb: None)
        with self.assertRaises(adapters.AdapterError):
            adapters.adapt("empty.xlsx", empty)
        with mock.patch.object(adapters, "MAX_XLSX_CELLS", 5), self.assertRaises(adapters.AdapterError):
            adapters.adapt("big.xlsx", raw)

    def test_macros_inside_a_renamed_file_are_refused(self):
        raw = xlsx_bytes(pricing_workbook)
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(buf, "w") as dst:
            for info in src.infolist():
                dst.writestr(info, src.read(info.filename))
            dst.writestr("xl/vbaProject.bin", b"\0")
        with self.assertRaises(adapters.AdapterError) as ctx:
            adapters.adapt("renamed.xlsx", buf.getvalue())
        self.assertIn("macros", str(ctx.exception))


def deck_bytes(build) -> bytes:
    deck = pptx.Presentation()
    build(deck)
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def proposal_deck(deck):
    s1 = deck.slides.add_slide(deck.slide_layouts[1])
    s1.shapes.title.text = "Harbour Bank: launch proposal"
    s1.placeholders[1].text = "Every NUSD payout on Polygon screened in real time before release, from the 1 December launch."
    s2 = deck.slides.add_slide(deck.slide_layouts[5])
    s2.shapes.title.text = "Commercials"
    table = s2.shapes.add_table(2, 3, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, row in enumerate([["Service", "Scope", "Effort"], ["Implementation", "Launch scope", "60 person-days"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    s2.notes_slide.notes_text_frame.text = "Do not promise 40,000/day yet."
    s3 = deck.slides.add_slide(deck.slide_layouts[5])
    s3.shapes.title.text = "Roadmap (internal)"
    s3._element.set("show", "0")


class TestPowerPoint(unittest.TestCase):
    def test_slides_in_order_title_first_tables_as_rows(self):
        out = adapters.adapt("proposal.pptx", deck_bytes(proposal_deck))
        self.assertEqual((out["format"], out["adapter_name"]), ("PowerPoint", "pptx"))
        text = out["canonical_text"]
        self.assertTrue(text.startswith("[Slide 1]\nHarbour Bank: launch proposal\nEvery NUSD payout on Polygon"))
        self.assertIn("[Slide 2]\nCommercials\nService | Scope | Effort\nImplementation | Launch scope | 60 person-days", text)
        self.assertIn("[Slide 3 (hidden)]\nRoadmap (internal)", text)
        self.assertNotIn("40,000/day", text)  # speaker notes are not read
        joined = " ".join(out["problems"])
        self.assertIn("Hidden slide 3", joined)
        self.assertIn("Speaker notes are not read", joined)

    def test_quotes_trace_to_their_slide(self):
        out = adapters.adapt("proposal.pptx", deck_bytes(proposal_deck))
        text, lmap = out["canonical_text"], out["location_map"]
        for e in json.loads(lmap):
            self.assertTrue(text[e["start"]:e["end"]].startswith(f"[Slide {e['slide']}"))
        self.assertEqual(adapters.locate(text, lmap, "Every NUSD payout on Polygon screened in real time", "pptx"), "slide 1")
        self.assertEqual(adapters.locate(text, lmap, "Roadmap (internal)", "pptx"), "slide 3 (hidden)")
        self.assertTrue(adapters.quote_in_text(text, "screened in real time before release, from the 1 December launch."))

    def test_groups_are_opened_and_pictures_noted(self):
        def build(deck):
            s = deck.slides.add_slide(deck.slide_layouts[6])
            group = s.shapes.add_group_shape()
            box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            box.text_frame.text = "Grouped promise text"
            from PIL import Image
            img = io.BytesIO()
            Image.new("RGB", (4, 4)).save(img, "PNG")
            img.seek(0)
            s.shapes.add_picture(img, Inches(4), Inches(4))
        out = adapters.adapt("g.pptx", deck_bytes(build))
        self.assertIn("Grouped promise text", out["canonical_text"])
        self.assertIn("1 picture, chart or object was not read", " ".join(out["problems"]))

    def test_refusals(self):
        raw = deck_bytes(proposal_deck)
        for name in ("old.ppt", "macro.pptm"):
            with self.assertRaises(adapters.AdapterError):
                adapters.adapt(name, raw)
        with self.assertRaises(adapters.AdapterError):
            adapters.adapt("empty.pptx", deck_bytes(lambda d: None))
        with mock.patch.object(adapters, "MAX_PPTX_SLIDES", 2), self.assertRaises(adapters.AdapterError):
            adapters.adapt("big.pptx", raw)
        with self.assertRaises(adapters.AdapterError):
            adapters.adapt("wrong.pptx", xlsx_bytes(pricing_workbook))  # a workbook renamed .pptx


class TestExcelPricingNoteEndToEnd(unittest.TestCase):
    """The 3 Oct decision: an Excel pricing note counts as approval evidence, by its content, like Markdown and Word."""

    def setUp(self):
        from test_recheck import build_current
        patcher = mock.patch.object(config, "LEDGER_DEALS", ["harbour_bank", "hard_cases"])
        patcher.start()
        self.addCleanup(patcher.stop)
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.db = Path(tmp.name) / "ledger.sqlite"
        build_current(self.db).close()
        self.client = TestClient(api.create_app(dist_dir=Path(tmp.name) / "none", ledger_path=self.db), base_url="http://127.0.0.1")

    def test_an_excel_named_exception_closes_the_approval_finding_only(self):
        import base64
        reg = self.client.get("/api/deals/harbour_bank/register").json()
        polygon = next(c for c in reg["commitments"] if "Polygon, real time" in c["name"])
        approval = next(i for i in polygon["issues"] if i["type"] == "approval")
        raw = xlsx_bytes(pricing_workbook)
        r = self.client.post("/api/deals/harbour_bank/findings/check", headers=H, json={
            "issue_id": approval["id"], "source_key": "HB-05", "confirmed_doc_type": "pricing_services_note",
            "filename": "HB-05_v2_pricing.xlsx", "file_base64": base64.b64encode(raw).decode(), "signed_off_by": "Daniel Koh"})
        self.assertEqual(r.status_code, 200, r.text)
        out = r.json()["check"]
        self.assertEqual(out["closed"], ["No approval recorded"])
        self.assertEqual(out["still_open"], ["Missing from contract", "Contract says something different"])
        docs = self.client.get("/api/deals/harbour_bank/documents").json()["documents"]
        self.assertEqual(next(d for d in docs if d["source_key"] == "HB-05")["format"], "Excel")

    def test_a_formula_approval_cell_with_no_stored_value_needs_evidence_and_never_closes(self):
        """Lina's decision (9 Oct, confirming 3 Oct): a missing formula value gives Needs evidence, never approval."""
        import base64
        import rules
        self.assertEqual(rules.NO_STORED_VALUE, adapters.NO_VALUE)

        def build(wb):
            pricing_workbook(wb)
            ws = wb.active
            for row in ws.iter_rows():
                if row[0].value and "Polygon, real-time" in str(row[1].value or ""):
                    row[3].value = "=Approvals!Z99"  # a formula Excel never calculated: no stored value
        reg = self.client.get("/api/deals/harbour_bank/register").json()
        polygon = next(c for c in reg["commitments"] if "Polygon, real time" in c["name"])
        approval = next(i for i in polygon["issues"] if i["type"] == "approval")
        r = self.client.post("/api/deals/harbour_bank/findings/check", headers=H, json={
            "issue_id": approval["id"], "source_key": "HB-05", "confirmed_doc_type": "pricing_services_note",
            "filename": "HB-05_v2_pricing.xlsx", "file_base64": base64.b64encode(xlsx_bytes(build)).decode(),
            "signed_off_by": "Daniel Koh"})
        self.assertEqual(r.status_code, 200, r.text)
        out = r.json()["check"]
        self.assertNotIn("No approval recorded", out["closed"])
        reg = self.client.get("/api/deals/harbour_bank/register").json()
        polygon = next(c for c in reg["commitments"] if "Polygon, real time" in c["name"])
        self.assertEqual(polygon["authorisation"], "Can't be confirmed")
        self.assertIn("Needs evidence", [i["finding"] for i in polygon["issues"] if i["state"] != "Resolved"])
        self.assertIn("formula with no saved value", polygon["authorisation_evidence"])
        self.assertNotIn("no approver", polygon["authorisation_evidence"])


if __name__ == "__main__":
    unittest.main()
