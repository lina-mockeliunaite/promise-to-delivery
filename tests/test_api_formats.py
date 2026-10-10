"""Uploading PDF and Word through the workspace API: formats, problems, refusals, page citations, pricing notes as evidence,
and the seal tests on the new upload path. Temporary ledger; a scripted fake stands in for the model. Nothing under
data/ is written and nothing under data/sealed_decoy/ is read (an audit hook checks the upload and review path).
"""

import base64
import csv
import io
import os
import re
import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent / "fixtures"))

import adapters
import config
import make_formats as mf
from test_api_workspace import H, WorkspaceCase
from test_recheck import FakeExtractor

BASE = "/api/deals/harbour_bank"
SENTINEL_KEY = "sk-ant-sentinel-do-not-leak-0123456789"
_active, _events = False, []


def _hook(event, args):
    if _active and event in {"open", "os.listdir", "os.scandir", "glob.glob"}:
        _events.append((event, tuple(str(a) for a in args)))


sys.addaudithook(_hook)


def b64(raw: bytes) -> str:
    return base64.b64encode(raw).decode("ascii")


def scenario_text() -> str:
    return config.scenario_path("HB-05_v2_named_exception.md").read_text(encoding="utf-8")


def md_lines(md: str) -> list:
    """The note's lines as text-layer lines: table separators dropped, markdown marks removed, no wrapping."""
    out = []
    for line in md.splitlines():
        line = line.strip()
        if not line or re.fullmatch(r"\|[-| ]+\|", line):
            continue
        out.append(re.sub(r"^#+\s*", "", line).replace("**", ""))
    return out


def md_docx(md: str) -> bytes:
    import docx
    document = docx.Document()
    rows = []

    def flush():
        if rows:
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, cells in enumerate(rows):
                for c, text in enumerate(cells):
                    table.rows[r].cells[c].text = text
            rows.clear()
    for line in md.splitlines():
        line = line.strip()
        if line.startswith("|"):
            if not re.fullmatch(r"\|[-| ]+\|", line):
                rows.append([c.strip() for c in line.strip("|").split("|")])
            continue
        flush()
        if line.startswith("#"):
            document.add_heading(line.lstrip("# ").strip(), level=2)
        elif line:
            document.add_paragraph(line.replace("**", ""))
    flush()
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class FormatsCase(WorkspaceCase):
    def upload(self, deal, filename, raw, **extra):
        return self.post(f"/api/deals/{deal}/documents", {"filename": filename, "file_base64": b64(raw), **extra})

    def user_deal(self):
        return self.post("/api/user-deals", {"name": "Formats deal"}).json()["deal"]

    def documents(self, deal="harbour_bank"):
        return self.client.get(f"/api/deals/{deal}/documents").json()["documents"]


class TestUploads(FormatsCase):
    def test_pdf_and_word_upload_as_new_documents_with_their_format_and_problems(self):
        deal = self.user_deal()
        r = self.upload(deal, "proposal.pdf", mf.hb04_pdf(), doc_type="proposal", doc_date="2026-10-10", name="x")
        self.assertEqual((r.status_code, r.json()["format"], r.json()["problems"]), (200, "PDF", []), r.text)
        r = self.upload(deal, "sow.docx", mf.hb04_docx(with_extras=True), doc_type="draft_sow", doc_date="2026-10-20")
        self.assertEqual((r.status_code, r.json()["format"]), (200, "Word"), r.text)
        self.assertEqual(r.json()["problems"], ["1 image was not read (pictures are not supported in this build)."])
        r = self.post(f"/api/deals/{deal}/documents", {"name": "n", "filename": "notes.txt", "doc_type": "customer_email",
                                                       "doc_date": "2026-10-27", "text": "A note.\n"})
        self.assertEqual(r.status_code, 200, r.text)
        docs = {d["doc_type_label"]: d for d in self.documents(deal)}
        self.assertEqual((docs["Proposal"]["format"], docs["Draft SOW"]["format"], docs["Customer email"]["format"]),
                         ("PDF", "Word", "Text"))
        self.assertEqual(docs["Draft SOW"]["problems"], ["1 image was not read (pictures are not supported in this build)."])
        self.assertEqual(docs["Proposal"]["problems"], [])
        self.assertTrue(all("filename" not in d for d in docs.values()))

    def test_a_scanned_page_is_reported_in_the_documents_list(self):
        deal = self.user_deal()
        r = self.upload(deal, "scan.pdf", mf.partial_scan_pdf(), doc_type="proposal", doc_date="2026-10-10")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(self.documents(deal)[0]["problems"],
                         ["Page 2 has no readable text (scanned or image-only; OCR is not supported in this build)."])

    def test_a_pdf_can_be_a_new_version_of_an_existing_source(self):
        before = {d["source_key"]: d for d in self.documents()}
        self.assertTrue(all(d["format"] == "Markdown" for d in before.values()))
        r = self.upload("harbour_bank", "HB-04.pdf", mf.hb04_pdf(), source_key="HB-04")
        self.assertEqual((r.status_code, r.json()["format"]), (200, "PDF"), r.text)
        now = {d["source_key"]: d for d in self.documents()}
        self.assertEqual((now["HB-04"]["version"], now["HB-04"]["format"]), (2, "PDF"))
        self.assertEqual(self.client.get(f"{BASE}/register").json()["freshness"]["state"], "Review out of date")

    def test_markdown_uploads_are_unchanged(self):
        r = self.post(f"{BASE}/documents", {"source_key": "HB-04", "filename": "HB-04.md", "text": "# Proposal\n\nA promise.\n"})
        self.assertEqual((r.status_code, r.json()["format"], r.json()["problems"]), (200, "Markdown", []))

    def test_a_database_that_never_had_a_document_with_problems_still_lists_documents(self):
        import ledger
        conn = ledger.connect(self.db)
        self.assertIsNone(conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'source_version_notes'").fetchone())
        conn.close()
        self.assertEqual(len(self.documents()), 8)
        self.upload(self.user_deal(), "scan.pdf", mf.partial_scan_pdf(), doc_type="proposal", doc_date="2026-10-10")
        conn = ledger.connect(self.db)
        self.assertIsNotNone(conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'source_version_notes'").fetchone())
        self.assertGreater(conn.execute("SELECT COUNT(*) FROM commitments").fetchone()[0], 0)  # existing data kept
        conn.close()


class TestRefusals(FormatsCase):
    def refused(self, filename, raw, status, fragment, deal="harbour_bank", **extra):
        r = self.upload(deal, filename, raw, **extra)
        self.assertEqual(r.status_code, status, r.text)
        self.assertIn(fragment, r.json()["detail"])
        self.assertNotIn(filename, r.json()["detail"])  # a file name is never echoed back
        return r

    def test_unreadable_and_unsupported_files_are_refused_with_a_plain_message(self):
        deal = self.user_deal()
        extra = {"doc_type": "proposal", "doc_date": "2026-10-10"}
        self.refused("scan.pdf", mf.image_only_pdf(), 400, "No page in this PDF has readable text", deal, **extra)
        self.refused("fake.pdf", b"not a pdf", 400, "not a PDF", deal, **extra)
        self.refused("old.doc", adapters.OLE_MAGIC + b"\0" * 100, 400, "Save the document as .docx", deal, **extra)
        self.refused("../../etc/passwd.exe", b"x", 400, "reads PDF, Word", deal, **extra)
        self.refused("archive.zip", b"PK\x03\x04", 400, "reads PDF, Word", deal, **extra)
        self.assertEqual(self.documents(deal), [])  # nothing was stored

    def test_oversize_files_are_refused(self):
        deal = self.user_deal()
        extra = {"doc_type": "proposal", "doc_date": "2026-10-10"}
        self.refused("big.pdf", b"%PDF-" + b"0" * adapters.MAX_UPLOAD_BYTES, 413, "larger than 10 MB", deal, **extra)
        huge = {"filename": "huge.pdf", "file_base64": "A" * (15 * 1024 * 1024), **extra}
        r = self.post(f"/api/deals/{deal}/documents", huge)  # announced size over the body limit: refused before reading
        self.assertEqual((r.status_code, r.json()["detail"]), (413, "The file is larger than 10 MB."))
        self.assertEqual(self.documents(deal), [])

    def test_bad_base64_a_missing_name_and_a_non_string_body_are_refused(self):
        deal = self.user_deal()
        r = self.post(f"/api/deals/{deal}/documents", {"filename": "a.pdf", "file_base64": "***not base64***",
                                                       "doc_type": "proposal", "doc_date": "2026-10-10"})
        self.assertEqual((r.status_code, r.json()["detail"]), (400, "The file could not be read. Upload it again."))
        r = self.post(f"/api/deals/{deal}/documents", {"file_base64": b64(mf.hb04_pdf()), "doc_type": "proposal", "doc_date": "2026-10-10"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("with its name", r.json()["detail"])
        r = self.post(f"/api/deals/{deal}/documents", {"filename": "a.pdf", "file_base64": 12345})
        self.assertEqual(r.status_code, 413)


class TestSeal(FormatsCase):
    def test_non_allowlisted_deals_are_404_before_the_file_is_read_and_writes_need_the_guard(self):
        body = {"filename": "a.pdf", "file_base64": b64(mf.hb04_pdf()), "doc_type": "proposal", "doc_date": "2026-10-10"}
        for deal in ("hard_cases", "sealed_decoy", "u_0123456789abcdef", "..", "harbour_bank%2F..%2Fsealed_decoy"):
            r = self.post(f"/api/deals/{deal}/documents", body)
            self.assertIn(r.status_code, (404, 422), deal)
            r = self.post(f"/api/deals/{deal}/documents", {**body, "file_base64": "A" * (15 * 1024 * 1024)})
            self.assertIn(r.status_code, (404, 422), deal)  # an oversize body to a deal we do not serve is still a 404
        self.assertEqual(self.client.post(f"{BASE}/documents", json=body).status_code, 403)
        r = self.client.post(f"{BASE}/documents", data="x", headers={**H, "Content-Type": "text/plain"})
        self.assertEqual(r.status_code, 403)
        self.assertEqual(len(self.documents()), 8)

    def test_the_upload_and_review_path_never_reads_sealed_decoy_or_writes_to_data(self):
        global _active
        deal = self.user_deal()
        payload = {"statements": [{"quote": "The payout ledger connector will handle up to 12,000 payouts per day at launch.",
                                   "speaker": "x", "language": "firm"}]}
        _events.clear()
        _active = True
        try:
            with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": SENTINEL_KEY}), \
                    mock.patch("anthropic.Anthropic", return_value=FakeExtractor(payload)):
                self.assertEqual(self.upload(deal, "p.pdf", mf.hb04_pdf(), doc_type="proposal", doc_date="2026-10-10").status_code, 200)
                self.assertEqual(self.upload(deal, "s.docx", mf.hb04_docx(), doc_type="draft_sow", doc_date="2026-10-20").status_code, 200)
                self.assertEqual(self.post(f"/api/deals/{deal}/review", {}).status_code, 200)
        finally:
            _active = False
        self.assertEqual([e for e in _events if any("sealed_decoy" in a.lower() for a in e[1])], [])
        data_root = str(config.DATA_DIR.resolve())
        self.assertEqual([e for e in _events if e[0] == "open" and len(e[1]) > 1 and e[1][0].startswith(data_root)
                          and any(m in e[1][1] for m in "wax+")], [])

    def test_no_key_path_or_file_name_reaches_any_response(self):
        deal = self.user_deal()
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": SENTINEL_KEY}):
            r = self.upload(deal, "Secret Name.pdf", mf.hb04_pdf(), doc_type="proposal", doc_date="2026-10-10")
        seen = r.text + str(self.documents(deal)) + self.client.get(f"/api/deals/{deal}/register").text
        for needle in (SENTINEL_KEY, "Secret Name", "/Users/", ".sqlite", "base64"):
            self.assertNotIn(needle, seen)


class TestQuotesAndPages(FormatsCase):
    def test_a_review_of_a_pdf_cites_the_page_in_the_register_and_in_both_exports(self):
        deal = self.user_deal()
        pdf = mf.text_pdf(mf.layout(mf.hb04_blocks(), 12))
        self.assertEqual(self.upload(deal, "p.pdf", pdf, doc_type="proposal", doc_date="2026-10-10").status_code, 200)
        quote = "The payout ledger connector will handle up to 12,000 payouts per day at launch."
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": "test-only"}), \
                mock.patch("anthropic.Anthropic", return_value=FakeExtractor({"statements": [
                    {"quote": quote, "speaker": "x", "language": "firm"}]})):
            self.assertEqual(self.post(f"/api/deals/{deal}/review", {}).status_code, 200)
        reg = self.client.get(f"/api/deals/{deal}/register").json()
        statement = reg["commitments"][0]["statements"][0]
        self.assertEqual(statement["quote"], quote)
        import json
        pages = [e["page"] for e in json.loads(adapters.adapt("p.pdf", pdf)["location_map"])]
        self.assertIn(statement["page"], pages)
        expected = next(p for p in pages if adapters.locate_page(adapters.adapt("p.pdf", pdf)["canonical_text"],
                                                                 adapters.adapt("p.pdf", pdf)["location_map"], quote) == p)
        self.assertEqual(statement["page"], expected)
        saved = self.post(f"/api/deals/{deal}/handoffs", {"decision": "not_ready", "reviewer": "Lina",
                                                           "review_id": reg["freshness"]["review_id"]})
        self.assertEqual(saved.status_code, 200, saved.text)
        rows = list(csv.reader(io.StringIO(self.client.get(f"/api/deals/{deal}/handoffs/export.csv").content.decode("utf-8-sig"))))
        self.assertTrue(any(f"(page {expected})" in r[6] for r in rows[1:]), rows)
        html = self.client.get(f"/api/deals/{deal}/handoffs/summary").text
        self.assertIn(f"version 1, page {expected}", html)

    def test_a_quote_from_a_docx_has_no_page(self):
        deal = self.user_deal()
        self.upload(deal, "p.docx", mf.hb04_docx(), doc_type="proposal", doc_date="2026-10-10")
        with mock.patch.dict(os.environ, {"ANTHROPIC_API_KEY": "test-only"}), \
                mock.patch("anthropic.Anthropic", return_value=FakeExtractor({"statements": [
                    {"quote": "NUSD payouts on Ethereum will be screened in real time through Elva's analytics integration before release.",
                     "speaker": "x", "language": "firm"}]})):
            self.assertEqual(self.post(f"/api/deals/{deal}/review", {}).status_code, 200)
        statement = self.client.get(f"/api/deals/{deal}/register").json()["commitments"][0]["statements"][0]
        self.assertIsNone(statement["page"])


class TestPricingNotesAsApprovalEvidence(FormatsCase):
    def close_approval_with(self, filename, raw):
        r = self.upload("harbour_bank", filename, raw, source_key="HB-05")
        self.assertEqual(r.status_code, 200, r.text)
        vid = r.json()["source_version_id"]
        polygon = next(c for c in self.register()["commitments"] if "Polygon, real time" in c["name"])
        approval = next(i for i in polygon["issues"] if i["type"] == "approval")
        fix = self.post(f"{BASE}/fixes", {"route": "allowed_exception", "owner": "Product", "rationale": "Named exception in the note.",
                                          "issue_ids": [approval["id"]], "approved_by": "Daniel Koh",
                                          "evidence": [{"source_version_id": vid, "locator": "whole document"}]})
        self.assertEqual(fix.status_code, 200, fix.text)
        r = self.post(f"{BASE}/review", {"fix_id": fix.json()["fix_id"]})
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["review"]["model_calls"], 0)
        polygon = next(c for c in self.register()["commitments"] if "Polygon, real time" in c["name"])
        return {i["type"]: i["state"] for i in polygon["issues"]}, polygon

    def test_a_word_pricing_note_counts_as_approval_evidence(self):
        states, polygon = self.close_approval_with("note.docx", md_docx(scenario_text()))
        self.assertEqual(states["approval"], "Resolved")
        self.assertEqual(states["contract_gap"], "Needs action")
        self.assertIn("Pricing note: named exception approved by Daniel Koh", polygon["authorisation_evidence"])
        self.assertEqual(next(d for d in self.documents() if d["source_key"] == "HB-05")["format"], "Word")

    def test_a_pdf_pricing_note_counts_as_approval_evidence(self):
        states, polygon = self.close_approval_with("note.pdf", mf.text_pdf([md_lines(scenario_text())]))
        self.assertEqual(states["approval"], "Resolved")
        self.assertIn("Pricing note: named exception approved by Daniel Koh", polygon["authorisation_evidence"])
        self.assertEqual(next(d for d in self.documents() if d["source_key"] == "HB-05")["format"], "PDF")


if __name__ == "__main__":
    unittest.main()
